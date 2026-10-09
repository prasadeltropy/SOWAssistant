#!/usr/bin/env python3
"""
SOW extraction + rendering engine for the Eltropy SOW-generator skill.

Given a CPQ workbook (.xlsx export of the CPQ Google Sheet), a package column
(Essentials / Enterprise / Eltropy One) and products.json (feature mapping),
produce:
  * sow_data.json  -- structured data (customer, selected SKUs, usage, pricing)
  * sow.html       -- a fully-rendered SOW ready to upload to Google Drive as a
                      Google Doc (create_file, contentMimeType text/html).

The HTML output is a FAITHFUL reproduction of the Eltropy SOW template
(doc 1x-NpQhL1T2Vdd8LlmIcFqfeZbW6oOijvXDsr0K3LQ1A): same sections, order and
boilerplate wording. Only the dynamic parts (customer name, Proposed Solutions /
Integrations rows, Phase 1 SKU lists) are filled from the CPQ.

Usage:
  python3 extract_sow.py --xlsx cpq.xlsx --column enterprise \
      --products products.json --out-json sow_data.json --out-html sow.html
"""
import argparse, json, re, html, sys
import openpyxl

PACKAGES = {
    "essentials": ("E", "Essentials"),
    "enterprise": ("F", "Enterprise"),
    "eltropy_one": ("G", "Eltropy One"),
    "eltropyone": ("G", "Eltropy One"),
    "e": ("E", "Essentials"), "f": ("F", "Enterprise"), "g": ("G", "Eltropy One"),
}

SECTION_LABELS = ["Platform", "Solutions", "Integrations", "Bundles & Misc"]
USAGE_LABEL = "Included Limited Usage"
PRICING_MARKERS = ["List Price", "Discount", "Discounted Price"]
PLACEHOLDERS = {"solutions", "integrations", "bundles & misc", "platform",
                "please click on", ""}


def norm(s):
    if s is None:
        return ""
    return re.sub(r"[^a-z0-9]+", " ", str(s).lower()).strip()


def is_included(v):
    if v is None:
        return False
    s = str(v).strip().lower()
    return s in ("included", "yes", "y", "true", "✓", "x")


def load_products(path):
    data = json.load(open(path))
    index = []  # (normalized_key, entry, category)
    for cat, entries in data["categories"].items():
        for e in entries:
            keys = [e["name"]] + e.get("aliases", [])
            for k in keys:
                index.append((norm(k), e, cat))
    index.sort(key=lambda t: len(t[0]), reverse=True)  # longest keys first
    return data, index


def match_product(name, index):
    n = norm(name)
    if not n:
        return None, None
    for key, e, cat in index:          # exact
        if key == n:
            return e, cat
    for key, e, cat in index:          # substring, either direction
        if key and (key in n or n in key):
            return e, cat
    return None, None


def find_package_header(ws):
    # Package column labels vary between CPQs (e.g. F may read "Enterprise Payments"),
    # so anchor on the Essentials (E) column and confirm the Eltropy One (G) column.
    for r in range(1, min(ws.max_row, 60) + 1):
        vals = {c: str(ws[c + str(r)].value or "").strip() for c in ("E", "F", "G")}
        e, g = vals["E"].lower(), vals["G"].lower()
        if e.startswith("essentials") and ("eltropy one" in g or "eltropy 1" in g):
            return r, vals
    # looser fallback: any row whose E is exactly 'Essentials'
    for r in range(1, min(ws.max_row, 60) + 1):
        if str(ws["E" + str(r)].value or "").strip().lower() == "essentials":
            return r, {c: str(ws[c + str(r)].value or "").strip() for c in ("E", "F", "G")}
    raise RuntimeError("Could not locate the Essentials / Enterprise / Eltropy One header "
                       "row (columns E/F/G, around row 15) in the Quote tab.")


def get_customer(ws):
    def cell(c):
        v = ws[c].value
        return "" if v is None else str(v).strip()
    return {
        "name": cell("C2") or cell("E2"),
        "full_name": cell("C3"),
        "aum": cell("C4"),
        "aum_bucket": cell("C5"),
        "members": cell("C6"),
        "employees": cell("C7"),
        "branches": cell("C8"),
    }


def apply_section_filters(sections, rules):
    """Honor `section_include_only` rules: within the named CPQ section, keep only
    the SKUs whose name matches one of the `keep` tokens (drop the rest)."""
    for rule in rules or []:
        if rule.get("type") != "section_include_only":
            continue
        sec = rule.get("section")
        keep = [norm(k) for k in rule.get("keep", [])]
        if sec not in sections or not keep:
            continue
        newlist = []
        for it in sections[sec]:
            nm = norm(it.get("matched_name") or it["name"])
            if any(k and (k in nm or nm in k) for k in keep):
                newlist.append(it)
        sections[sec] = newlist
    return sections


def extract(ws, col, index):
    hdr_row, _ = find_package_header(ws)
    sections = {s: [] for s in SECTION_LABELS}
    usage = []
    pricing = {}
    current = None
    mode = "items"
    r = hdr_row + 1
    while r <= ws.max_row:
        b = str(ws["B" + str(r)].value or "").strip()
        c_name = str(ws["C" + str(r)].value or "").strip()
        a_upi = str(ws["A" + str(r)].value or "").strip()
        val = ws[col + str(r)].value

        if b in SECTION_LABELS:
            current = b
            mode = "items"
        elif b == USAGE_LABEL:
            mode = "usage"
            current = None
        elif any(m.lower() in b.lower() for m in PRICING_MARKERS) and b:
            mode = "pricing"

        if mode == "pricing":
            label = b or c_name
            if label:
                pricing[label] = "" if val is None else str(val).strip()
            r += 1
            if r > hdr_row + 200:
                break
            continue

        if mode == "usage":
            if c_name and norm(c_name) not in PLACEHOLDERS:
                sval = "" if val is None else str(val).strip()
                if sval and sval not in ("0", "0.0", "-"):
                    usage.append({"name": c_name, "value": sval})
            r += 1
            continue

        if current and is_included(val) and norm(c_name) not in PLACEHOLDERS:
            entry, cat = match_product(c_name, index)
            sections.setdefault(current, []).append({
                "upi": a_upi,
                "name": c_name,
                "matched_name": entry["name"] if entry else None,
                "category": cat if entry else current,
                "cpq_section": current,
                "features": entry["features"] if entry else [],
                "notes": entry.get("notes", []) if entry else [],
                "first_value_moment": entry.get("first_value_moment") if entry else None,
                "acceptance_criteria": entry.get("acceptance_criteria") if entry else None,
            })
        r += 1
    return {"sections": sections, "usage": usage, "pricing": pricing}


# =====================================================================
# HTML rendering -- faithful to the Eltropy SOW template
# =====================================================================

TECH_REQS = [
    ("Computer and processor", "Minimum 1.1 GHz or faster, two core"),
    ("Memory", "8.0 GB RAM"),
    ("Hard disk", "3.0 GB of available disk space"),
    ("Display", "1024 x 768 screen resolution"),
    ("Graphics hardware", "Windows OS: Graphics hardware acceleration requires DirectX 9 or later, with WDDM 2.0 or higher for Windows 10 (or WDDM 1.3 or higher for Windows 10 Fall Creators Update)"),
    ("Operating system", "Windows 11, Windows 10 (excluding Windows 10 LTSC for Teams desktop app), Windows 10 on ARM, Windows 8.1, Windows Server 2019, Windows Server 2016, Windows Server 2012 R2. Note: We recommend using the latest Windows version and security patches available."),
    (".NET version", "Requires .NET 4.5 CLR or later"),
    ("Video", "USB 2.0 video camera"),
    ("Devices", "Standard laptop camera, microphone, and speakers"),
    ("Video calls and meetings", None),  # rendered as bullets, see below
    ("Privacy Policy Updates", None),
]

VIDEO_CALLS_BULLETS = [
    "Requires a two-core processor. For higher video/screen share resolution and frame rate, a four-core processor or better is recommended.",
    "Background video effects require Windows 10 or a processor with AVX2 instruction set.",
]

NETWORK_REQS = [
    "High-Speed Internet Connection: We recommend 100 mbps and above.",
    "QoS-Enabled Router: Configure your router to prioritize voice traffic (Eltropy's traffic) over other network traffic.",
    "TCP Ports: 80 and 443 for agent browser sessions should be open on Firewall/Router",
    "UDP Ports: 5060, plus UDP ports in the range of 10000-35000.",
    "VLANs: If running a converged network for voice and data, configure VLANs to separate the traffic.",
    "DNS and NTP: Ensure that the Phone VLAN has the following DNS and NTP in its DHCP scope:",
    "Cabling: Use at least Cat 5 (preferably Cat 6) cabling to your phone and PC.",
    "Network Switches: Ensure your network switch can properly handle the number of devices if you are using both PC and phone on a single network run.",
    "ALG: Eltropy's recommends that ALG (Application Layer Gateway) be disabled on all your network gear, as it is known to cause the Eltropy's solution to not perform correctly",
]

TOC_SECTIONS = [
    "Solution Overview", "Scope of Work", "Client Business Problems and Desired Outcomes",
    "Project Overview", "Proposed Solutions", "Deployment Phases", "Phase 1",
    "Risks and Assumptions - Phase 1", "Phase 2", "Risks and Assumptions - Phase 2",
    "Phase 3", "Risks and Assumptions - Phase 3", "Out of Scope", "Change Process",
    "Technical/Network Requirements", "Minimum Technical Requirements",
    "Network Requirements", "Acknowledgement of Receipt",
]

FOOTNOTE = ("*All items described in this and other Sections (including the list of "
            "solutions, features, first value moment and full value timelines) are not "
            "contractual deliverables, acceptance criteria, or a representation that such "
            "items will be delivered outside the terms of the Order Form or Agreement. The "
            "First Value Moment and Full Value moment described above are based on a summary "
            "of benefits that deployed customers have experienced. These timelines and "
            "outcomes are dependent on Client usage, organizational readiness, and other "
            "factors.")

CHANGE_PROCESS = ("Any change to this SOW should be communicated via email to your Account "
                  "Executive or Customer Success Manager who will coordinate with the "
                  "appropriate Eltropy teams. The Account Executive or Customer Success "
                  "Manager will work with you to assess how the proposed changes may impact "
                  "project timelines and goals, and will provide guidance on the best path "
                  "forward. Changes that may affect the Services, timelines or other "
                  "commercial terms must be implemented through a new Order Form executed "
                  "between the parties.")

ACK = ("By signing below, Client acknowledges receipt of this SOW for informational and "
       "planning purposes. Rejection or cancellation of this SOW shall not terminate or "
       "otherwise affect any Order Form or the Agreement.")


def esc(s):
    return html.escape(str(s), quote=False)


def bullets(items):
    return "<ul>" + "".join(f"<li>{esc(i)}</li>" for i in items) + "</ul>"


def render_solution_rows(items):
    """One <tr> per product: name | features (+notes) | first value moment."""
    rows = ""
    for it in items:
        if it["features"]:
            feat = bullets(it["features"])
        else:
            feat = "<ul><li>[Insert from Product Marketing Descriptions]</li></ul>"
        if it["notes"]:
            feat += "<p>&nbsp;</p><p>Notes:</p>" + bullets(it["notes"])
        fvm = it.get("first_value_moment") or ""
        name = it["matched_name"] or it["name"]
        rows += f"<tr><td>{esc(name)}</td><td>{feat}</td><td>{esc(fvm)}</td></tr>"
    return rows


def render(data, extracted, package_label, customer, sow_date):
    sections = extracted["sections"]
    solutions = sections.get("Platform", []) + sections.get("Solutions", []) + \
        sections.get("Bundles & Misc", [])
    integrations = sections.get("Integrations", [])
    cust = esc(customer.get("name") or "[Client]")

    sol_names = [(s["matched_name"] or s["name"]) for s in solutions]
    int_names = [(s["matched_name"] or s["name"]) for s in integrations]

    P = []
    A = P.append
    A("<html><body>")

    # ---- Cover ----
    A(f"<h1>Eltropy for {cust}</h1>")
    A("<h2>Statement of Work Summary</h2>")
    A("<table border='1'><tr><th>Presented To</th><th>Presented By</th></tr>"
      f"<tr><td>{cust} (“Company” or “Client”)</td>"
      "<td>Eltropy, Inc. (“Eltropy”)<br/>Approved by: _______</td></tr></table>")
    A(f"<table border='1'><tr><th>Date</th><td>{esc(sow_date)}</td></tr></table>")

    # ---- Table of Contents (static list; see note in SKILL.md) ----
    A("<p><strong>Table of Contents</strong></p>")
    A("<ul>" + "".join(f"<li>{esc(s)}</li>" for s in TOC_SECTIONS) + "</ul>")

    A("<p><strong>IMPORTANT - NON-BINDING; FOR INFORMATIONAL PURPOSES ONLY</strong></p>")

    # ---- Solution Overview ----
    A("<h1>Solution Overview</h1>")
    A(f"<p>Eltropy is pleased to provide {cust} with this Statement of Work for {cust} "
      "(this “SOW”) which is provided solely for the Client’s planning and "
      "informational use to describe Eltropy’s anticipated implementation approach, "
      "phases and timelines. This SOW is non-binding, does not create any legal rights or "
      "obligations, and does not form part of, amend, or otherwise alter any Order Form, the "
      "Eltropy SaaS Agreement (the “Agreement”), or Client’s payment or other "
      "obligations thereunder.</p>")
    A(f"<p>Please review this document to ensure we have a proper understanding of {cust} "
      "goals for implementation of the Eltropy Services. We aim to ensure that both parties "
      "have comprehensive knowledge of the project's requirements, timeline and "
      "expectations.</p>")
    A("<p><strong>Kickoff:</strong></p>")
    A("<ul>"
      "<li>Upon signature of the Order Form, Eltropy will assign the following resources:"
      "<ul><li>Customer Success Manager (CSM) who will be the point contact moving forward.</li>"
      "<li>Implementation Manager/s (IM) who will be responsible for the ultimate success of "
      "the outlined project.</li></ul></li>"
      "<li>The Implementation Manager/s will prepare Implementation SOWs for each purchased "
      "solution that outline proposed deployment stages and estimated timelines based on "
      "stated requirements and Client resource availability.</li>"
      "<li>The Account Executive will set up introduction meetings with the above mentioned "
      "resources to plan the official project kickoff call.</li>"
      "<li>Post implementation, the CSM and IM will meet with you to review and evaluate the "
      "execution of the deployment, review and measure value realization.</li></ul>")

    # ---- Scope of Work ----
    A("<h1>Scope of Work</h1>")
    A("<h2>Client Business Problems and Desired Outcomes</h2>")
    A("<table border='1'><tr><th>Business Problems / Areas of Innovation</th><th>Outcomes</th></tr>"
      "<tr><td>&nbsp;</td><td>&nbsp;</td></tr><tr><td>&nbsp;</td><td>&nbsp;</td></tr>"
      "<tr><td>&nbsp;</td><td>&nbsp;</td></tr></table>")

    A("<h2>Project Overview</h2>")
    A("<p>This project involves the implementation of the proposed solutions and integration.</p>")

    A("<h2>Proposed Solutions</h2>")
    A("<p>Overview of the Solution described to the Client:</p>")
    A("<table border='1'><tr><th>Solutions</th><th>Features Included</th>"
      "<th>First Value Moment*</th></tr>")
    if solutions:
        A(render_solution_rows(solutions))
    else:
        A("<tr><td>[Solution]</td><td><ul><li>[Insert from Product Marketing Descriptions]"
          "</li></ul></td><td>[Key detail]</td></tr>")
    A("<tr><td colspan='3'>Proposed Integrations:</td></tr>")
    if integrations:
        A(render_solution_rows(integrations))
    else:
        A("<tr><td>[Integration App]</td><td><ul><li>[Insert from Product Marketing "
          "Descriptions]</li></ul></td><td>&nbsp;</td></tr>")
    A("<tr><td>Customization and/or anything else committed to the Client</td>"
      "<td>&nbsp;</td><td>[Key detail]</td></tr>")
    A("</table>")
    A(f"<p><em>{esc(FOOTNOTE)}</em></p>")

    # ---- Deployment Phases ----
    A("<h2>Deployment Phases</h2>")
    A(f"<p>Eltropy and {cust} will work to implement in [x] phases. The estimated deployment "
      "timeline may vary based on Client readiness and resource availability. Each solution "
      "will have an Implementation SOW, created by the Implementation Manager/s with your "
      "input on goals and success criteria and a specific timeline.</p>")

    def phase(num, sols, ints):
        A(f"<h3>Phase {num}</h3>")
        A("<ul><li>Expected Start Date  : </li><li>Expected End Date  : </li>"
          "<li>Project Manager Name : </li></ul>")
        A("<table border='1'><tr><th>SKU’s to be Implemented</th>"
          "<th>Client Resource Requirement</th></tr>")
        sol_cell = "Solutions:" + (bullets(sols) if sols else "<ul><li>[Insert item]</li></ul>")
        int_cell = "Integrations:" + (bullets(ints) if ints else "<ul><li>[Insert item]</li></ul>")
        A(f"<tr><td>{sol_cell}</td><td>&nbsp;</td></tr>")
        A(f"<tr><td>{int_cell}</td><td>&nbsp;</td></tr>")
        A("</table>")
        A(f"<h4>Risks and Assumptions - Phase {num}</h4><ul><li>&nbsp;</li></ul>")

    phase(1, sol_names, int_names)
    phase(2, None, None)
    phase(3, None, None)

    A("<h2>Out of Scope</h2><ul><li>[List out what is not included]</li></ul>")

    A("<h2>Change Process</h2>")
    A(f"<p>{esc(CHANGE_PROCESS)}</p>")

    # ---- Technical / Network Requirements ----
    A("<h1>Technical/Network Requirements</h1>")
    A("<h2>Minimum Technical Requirements</h2>")
    A("<p>The Company will ensure the following requirements are met, for each product "
      "offering.</p>")
    A("<table border='1'><tr><th>Component</th><th>Requirement</th></tr>")
    for comp, req in TECH_REQS:
        if comp == "Video calls and meetings":
            cell = bullets(VIDEO_CALLS_BULLETS)
        elif comp == "Privacy Policy Updates":
            cell = ("<ul><li>If purchasing the following SKUs, " + cust +
                    " may need to review and update their privacy policy: Text Messaging, "
                    "RCS, Video Banking, Voice+, Eltropy Identity Verification.</li></ul>")
        else:
            cell = esc(req)
        A(f"<tr><td>{esc(comp)}</td><td>{cell}</td></tr>")
    A("</table>")

    A("<h2>Network Requirements</h2>")
    A("<table border='1'><tr><td>Network Requirements</td>"
      f"<td>{bullets(NETWORK_REQS)}</td></tr></table>")

    # ---- Acknowledgement ----
    A("<h1>Acknowledgement of Receipt</h1>")
    A(f"<p><em>{esc(ACK)}</em></p>")
    A("<p>By &nbsp; &nbsp; &nbsp; &nbsp; &nbsp; &nbsp; &nbsp; &nbsp; &nbsp;<br/>"
      "Name: &nbsp; &nbsp; &nbsp; &nbsp; &nbsp; &nbsp; &nbsp; &nbsp;<br/>"
      "Title: &nbsp; &nbsp; &nbsp; &nbsp; &nbsp; &nbsp; &nbsp;<br/>"
      "Date: &nbsp; &nbsp; &nbsp; &nbsp; &nbsp;</p>")

    A("</body></html>")
    return "".join(P)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--column", required=True, help="essentials | enterprise | eltropy_one (or E/F/G)")
    ap.add_argument("--products", required=True)
    ap.add_argument("--rules", default=None, help="rules.json; applies section_include_only filters")
    ap.add_argument("--sheet", default="Quote")
    ap.add_argument("--customer", default=None, help="override customer name")
    ap.add_argument("--date", default=None)
    ap.add_argument("--out-json", default="sow_data.json")
    ap.add_argument("--out-html", default="sow.html")
    args = ap.parse_args()

    key = args.column.strip().lower().replace(" ", "_")
    if key not in PACKAGES:
        sys.exit(f"Unknown column '{args.column}'. Use essentials | enterprise | eltropy_one.")
    col, package_label = PACKAGES[key]

    data, index = load_products(args.products)
    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    ws = wb[args.sheet] if args.sheet in wb.sheetnames else wb["Quote"]

    customer = get_customer(ws)
    if args.customer:
        customer["name"] = args.customer
    import datetime
    sow_date = args.date or datetime.date.today().strftime("%b %d, %Y")

    extracted = extract(ws, col, index)
    if args.rules:
        try:
            rules = json.load(open(args.rules)).get("rules", [])
            apply_section_filters(extracted["sections"], rules)
        except Exception as e:
            print("WARN: could not apply rules:", e)
    out = {
        "package": package_label, "column": col,
        "customer": customer,
        "sections": extracted["sections"],
        "usage": extracted["usage"],
        "pricing": extracted["pricing"],
    }
    json.dump(out, open(args.out_json, "w"), indent=2)
    open(args.out_html, "w").write(render(data, extracted, package_label, customer, sow_date))

    n_sol = sum(len(extracted["sections"].get(s, [])) for s in ("Platform", "Solutions", "Bundles & Misc"))
    n_int = len(extracted["sections"].get("Integrations", []))
    unmatched = [i["name"] for s in extracted["sections"].values() for i in s if not i["matched_name"]]
    print(f"Package: {package_label} (col {col})")
    print(f"Customer: {customer.get('name')}")
    print(f"Selected solutions/bundles: {n_sol} | integrations: {n_int} | usage rows: {len(extracted['usage'])}")
    if unmatched:
        print("UNMATCHED (no product description found):", unmatched)
    print("Wrote", args.out_json, "and", args.out_html)


if __name__ == "__main__":
    main()
